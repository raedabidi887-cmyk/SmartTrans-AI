from docx import Document
from fpdf import FPDF
from pathlib import Path
import re

def write_translated_file(output_path: Path, translated_content: str, original_extension: str):
    """
    Écrit le contenu traduit dans un fichier selon l'extension originale
    
    Args:
        output_path (Path): Chemin du fichier de sortie
        translated_content (str): Contenu traduit
        original_extension (str): Extension du fichier original
    """
    try:
        if original_extension == '.pdf':
            write_pdf_file(output_path, translated_content)
        elif original_extension == '.docx':
            write_docx_file(output_path, translated_content)
        elif original_extension == '.txt':
            write_txt_file(output_path, translated_content)
        else:
            raise ValueError(f"Format de fichier non supporté: {original_extension}")
            
    except Exception as e:
        print(f"Erreur lors de l'écriture du fichier {output_path}: {e}")

def write_pdf_file(output_path: Path, content: str):
    """
    Crée un fichier PDF avec le contenu traduit
    
    Args:
        output_path (Path): Chemin du fichier PDF de sortie
        content (str): Contenu à écrire
    """
    try:
        pdf = FPDF()
        pdf.add_page()
        
        # Configuration de la police
        pdf.add_font('DejaVu', '', 'backend/utils/fonts/DejaVuSansCondensed.ttf', uni=True)
        pdf.set_font('DejaVu', '', 12)
        
        # Diviser le contenu en lignes
        lines = content.split('\n')
        
        for line in lines:
            # Diviser les longues lignes
            words = line.split()
            current_line = ""
            
            for word in words:
                if pdf.get_string_width(current_line + " " + word) < 190:  # Marge de 10mm
                    current_line += " " + word if current_line else word
                else:
                    if current_line:
                        pdf.cell(0, 10, current_line, ln=True)
                    current_line = word
            
            if current_line:
                pdf.cell(0, 10, current_line, ln=True)
        
        pdf.output(str(output_path))
        
    except Exception as e:
        print(f"Erreur lors de la création du PDF: {e}")
        # Fallback: créer un fichier texte
        write_txt_file(output_path.with_suffix('.txt'), content)

def write_docx_file(output_path: Path, content: str):
    """
    Crée un fichier DOCX avec le contenu traduit
    
    Args:
        output_path (Path): Chemin du fichier DOCX de sortie
        content (str): Contenu à écrire
    """
    try:
        doc = Document()
        
        # Diviser le contenu en paragraphes
        paragraphs = content.split('\n')
        
        for paragraph_text in paragraphs:
            if paragraph_text.strip():
                doc.add_paragraph(paragraph_text)
        
        doc.save(str(output_path))
        
    except Exception as e:
        print(f"Erreur lors de la création du DOCX: {e}")
        # Fallback: créer un fichier texte
        write_txt_file(output_path.with_suffix('.txt'), content)

def write_txt_file(output_path: Path, content: str):
    """
    Crée un fichier TXT avec le contenu traduit
    
    Args:
        output_path (Path): Chemin du fichier TXT de sortie
        content (str): Contenu à écrire
    """
    try:
        with open(output_path, 'w', encoding='utf-8') as file:
            file.write(content)
            
    except Exception as e:
        print(f"Erreur lors de la création du TXT: {e}")
        # Dernier recours: essayer avec une autre encodage
        try:
            with open(output_path, 'w', encoding='latin-1') as file:
                file.write(content)
        except Exception as e2:
            print(f"Erreur critique lors de l'écriture du fichier: {e2}") 